"""Configuration constants for the legal report web application."""

from pathlib import Path

# Project root
PROJECT_ROOT = Path(__file__).parent.parent.resolve()

# Existing scripts and output
LAW_RESEARCH_DIR = PROJECT_ROOT / ".law-research"
SCRIPTS_DIR = LAW_RESEARCH_DIR / "scripts"
LAW_OUTPUT_DIR = LAW_RESEARCH_DIR / "output"

# Web app output
WEBAPP_OUTPUT_DIR = Path(__file__).parent / "output"
WEBAPP_OUTPUT_DIR.mkdir(exist_ok=True)

# Retry configuration
MAX_RETRIES = 2

# LLM temperature — low for consistency
DEFAULT_TEMPERATURE = 0.3

# Max tokens per provider
MAX_OUTPUT_TOKENS = {
    "openai": 16000,
    "anthropic": 32000,
    "deepseek": 8000,
    "custom": 16000,
}

# Default models per provider
DEFAULT_MODELS = {
    "openai": "gpt-4o",
    "anthropic": "claude-sonnet-4-20250514",
    "deepseek": "deepseek-chat",
    "custom": "",
}

# DeepSeek base URL
DEEPSEEK_BASE_URL = "https://api.deepseek.com/v1"

# Cover template
COVER_TEMPLATE_PATH = PROJECT_ROOT / "封面.docx"


def _ensure_cover_template():
    """如果封面模板文件不存在，自动生成一个最简版本。

    生成的封面包含【标题】和日期占位符，与 _replace_cover_placeholders 兼容。
    当仓库不包含封面文件时（如 Render 部署），确保报告仍能正常生成封面。
    """
    if COVER_TEMPLATE_PATH.exists():
        return

    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # 页面设置
        section = doc.sections[0]

        # 空行留白
        for _ in range(6):
            doc.add_paragraph()

        # 标题占位符
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run("【标题】")
        run_title.font.size = Pt(22)
        run_title.bold = True

        # 副标题
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run("法律评估意见书")
        run_sub.font.size = Pt(16)

        # 空行
        doc.add_paragraph()
        doc.add_paragraph()

        # 日期占位符
        p_date = doc.add_paragraph()
        p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_date = p_date.add_run("【二○二六】年 【六】月")
        run_date.font.size = Pt(12)

        doc.save(str(COVER_TEMPLATE_PATH))
    except Exception:
        pass  # 生成失败不影响应用启动，报告将无封面


_ensure_cover_template()

# URL validation
URL_VALIDATION_ENABLED = True           # Master switch for URL reachability checks
URL_VALIDATION_TIMEOUT = 10             # Seconds per URL HEAD/GET request
URL_VALIDATION_MAX_CONCURRENT = 5       # Max concurrent URL checks
URL_VALIDATION_MAX_URLS = 50            # Cap total URLs checked per report
URL_SKIP_DOMAINS = [                    # Domains assumed stable (skip HEAD check)
    "gov.cn",
    "npc.gov.cn",
]

# Web search for legal case discovery
SEARCH_ENABLED = True            # Master switch; set False to disable all web search
SEARCH_MAX_RESULTS = 8           # Max total search results across all queries
SEARCH_TIMEOUT = 15              # Seconds per HTTP request
SEARCH_MAX_QUERIES = 2           # Max distinct search queries to send
SEARCH_QUERY_DELAY = 2.0         # Seconds delay between queries (rate limiting)

# File cleanup: keep last N reports
MAX_OUTPUT_FILES = 50
